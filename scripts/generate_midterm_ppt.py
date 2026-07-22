#!/usr/bin/env python3
"""Generate the editable Fudan-style Agentic RAG-VLM midterm presentation."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "中期" / "中期汇报PPT_陈涛.pptx"
RAG_FIG = ROOT / "paper" / "Agentic-RAG-VLM" / "figures"
AGENTIC_VLA_FIG = ROOT / "paper" / "CARVE-VLA" / "figures"
BRAND = ROOT / "assets" / "presentation" / "fudan"
MIDTERM_ASSET = ROOT / "assets" / "presentation" / "midterm"

FUDAN = RGBColor(31, 75, 148)
BLUE = RGBColor(52, 103, 184)
LIGHT_BLUE = RGBColor(231, 239, 251)
TEAL = RGBColor(25, 135, 125)
LIGHT_TEAL = RGBColor(232, 247, 244)
ORANGE = RGBColor(222, 120, 35)
LIGHT_ORANGE = RGBColor(252, 241, 229)
RED = RGBColor(183, 58, 53)
GREEN = RGBColor(54, 135, 83)
INK = RGBColor(22, 27, 34)
MUTED = RGBColor(86, 96, 108)
LINE = RGBColor(195, 202, 211)
PANEL = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
FONT = "Microsoft YaHei"


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03,
             font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for idx, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, line_width=0.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_rule(slide, x, y, w, color=LINE, height=0.018):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_image_contain(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
        width=Inches(pw), height=Inches(ph)
    )


def add_image_panel(slide, path, x, y, w, h, padding=0.08):
    add_box(slide, x, y, w, h, WHITE, RGBColor(219, 225, 233), 0.7)
    return add_image_contain(slide, path, x + padding, y + padding,
                             w - 2 * padding, h - 2 * padding)


def add_logo(slide):
    add_image_contain(slide, BRAND / "fudan_logo_blue.png", 10.64, 0.17, 2.10, 0.67)


def add_header(slide, title, page, section=None):
    if section:
        add_text(slide, section, 0.52, 0.16, 2.7, 0.22, 8.5, FUDAN, True)
    add_text(slide, title, 0.52, 0.43, 9.72, 0.46, 24, INK, True)
    add_logo(slide)
    add_rule(slide, 0.50, 1.02, 12.32, LINE, 0.016)
    add_text(slide, f"{page:02d}", 12.18, 7.12, 0.55, 0.18, 8.5, MUTED,
             align=PP_ALIGN.RIGHT)


def add_footer(slide, text="具身多模态智能体系统研究进展"):
    return None


def remove_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[index])


def renumber_slides(prs):
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            left = shape.left / Inches(1)
            top = shape.top / Inches(1)
            width = shape.width / Inches(1)
            text = shape.text.strip()
            if left > 11.8 and top > 6.8 and width < 1.0 and text.isdigit():
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                run = p.add_run()
                run.text = f"{idx:02d}"
                run.font.name = FONT
                run.font.size = Pt(8.5)
                run.font.color.rgb = MUTED


def add_bullets(slide, items, x, y, w, h, size=15, accent=FUDAN, spacing=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.space_after = Pt(spacing)
        p.line_spacing = 1.08
        dot = p.add_run()
        dot.text = "●  "
        dot.font.name = FONT
        dot.font.size = Pt(max(8, size - 4))
        dot.font.color.rgb = accent
        run = p.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return box


def add_label(slide, text, x, y, w, color=FUDAN):
    add_box(slide, x, y, w, 0.34, color, color, 0)
    add_text(slide, text, x, y + 0.015, w, 0.26, 10.5, WHITE, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_section_tag(slide, text, x, y, w, color=FUDAN):
    add_text(slide, text, x, y, w, 0.22, 9.5, color, True,
             align=PP_ALIGN.CENTER)
    add_rule(slide, x, y + 0.28, w, color, 0.025)


def add_stat(slide, value, label, x, y, w, color=FUDAN):
    add_text(slide, value, x, y, w, 0.54, 28, color, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y + 0.57, w, 0.46, 11.5, MUTED, True,
             align=PP_ALIGN.CENTER)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover: same institutional visual language as the Fudan reference deck.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_image_contain(slide, BRAND / "fudan_campus_cover.png", 0.42, 0.0, 3.72, 7.50)
    add_logo(slide)
    add_text(slide, "硕士中期汇报", 4.60, 1.65, 7.55, 0.70, 31, INK, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "面向具身智能的多模态智能体系统与高效推理研究",
             4.45, 2.52, 7.85, 0.55, 20, FUDAN, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 5.58, 3.51, 5.48, 0.46, FUDAN, FUDAN, 0)
    add_text(slide, "汇报人：陈涛    导师：甘中学、刘力政", 5.62, 3.57, 5.40, 0.30,
             13, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "学号：24210860031", 4.65, 4.44, 7.47, 0.31, 13, MUTED,
             align=PP_ALIGN.CENTER)
    add_text(slide, "复旦大学 智能机器人与先进制造创新学院",
             4.65, 4.87, 7.47, 0.31, 13, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "2026 年 7 月 27 日", 4.65, 5.31, 7.47, 0.31, 13, MUTED,
             align=PP_ALIGN.CENTER)

    # 2. Agenda + research route.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "汇报结构与研究路线", 2)
    add_footer(slide)
    agenda = [
        ("01", "已有成果", "Agentic RAG-VLM", "IROS 2026录用；经验检索、场景约束与失败恢复", FUDAN),
        ("02", "近期进展", "Agentic VLA", "Agentic Harness 与 Optimize Runtime", ORANGE),
        ("03", "后续计划", "系统完善与验证", "完善闭环执行、运行时管理与实验验证", TEAL),
    ]
    for i, (num, sec, title, desc, color) in enumerate(agenda):
        y = 1.34 + i * 0.74
        add_text(slide, num, 0.76, y, 0.64, 0.36, 18, color, True)
        add_rule(slide, 1.44, y + 0.14, 0.62, color, 0.035)
        add_text(slide, sec, 2.20, y + 0.01, 1.30, 0.28, 13.5, INK, True)
        add_text(slide, title, 3.55, y + 0.01, 2.75, 0.28, 13.5, color, True)
        add_text(slide, desc, 6.25, y + 0.02, 4.70, 0.26, 11.5, MUTED)

    route = [
        ("结构化操作推理", "Agentic RAG-VLM", "VLM 外层：经验检索、场景约束、反思记忆", FUDAN),
        ("闭环执行增强", "Agentic VLA", "冻结 VLA 外层：监测、恢复、安全停止", ORANGE),
        ("运行时效率优化", "Optimize Runtime", "管理调用预算、动作复用和运行轨迹", TEAL),
    ]
    for i, (role, name, body, color) in enumerate(route):
        x = 0.68 + i * 4.18
        add_box(slide, x, 4.08, 3.58, 1.35, WHITE, RGBColor(210, 218, 228), 0.8)
        add_rule(slide, x, 4.08, 3.58, color, 0.06)
        add_text(slide, role, x + 0.22, 4.37, 3.10, 0.26, 13.2, MUTED, True)
        add_text(slide, name, x + 0.22, 4.73, 3.10, 0.32, 15.5, color, True)
        add_text(slide, body, x + 0.22, 5.11, 3.10, 0.32, 10.8, INK)
        if i < 2:
            add_text(slide, "→", x + 3.70, 4.75, 0.34, 0.42, 18, LINE, True,
                     align=PP_ALIGN.CENTER)
    add_box(slide, 0.68, 5.88, 11.96, 0.53, LIGHT_BLUE, RGBColor(182, 204, 235))
    add_text(slide, "研究目标：构建可解释、可恢复、可部署的具身多模态智能体系统。",
             0.90, 6.00, 11.50, 0.36, 13.5, FUDAN, True,
             align=PP_ALIGN.CENTER)

    # 3. Overall research route.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "整体研究路线", 3, "研究主线")
    add_footer(slide)
    stages = [
        ("已完成", "结构化操作推理", "Agentic RAG-VLM", "检索操作经验\n施加场景约束\n反思与记忆", FUDAN),
        ("近期进展", "闭环执行增强", "Agentic VLA", "监测执行状态\n选择恢复动作\n验证与安全停止", ORANGE),
        ("近期进展", "运行时效率优化", "Optimize Runtime", "管理调用预算\n复用动作块\n记录运行轨迹", TEAL),
    ]
    for i, (status, role, name, body, color) in enumerate(stages):
        x = 0.66 + i * 4.18
        add_label(slide, status, x, 1.55, 1.18, color)
        add_text(slide, role, x, 2.07, 3.46, 0.38, 16, MUTED, True)
        add_text(slide, name, x, 2.56, 3.48, 0.54, 21, color, True)
        add_rule(slide, x, 3.21, 3.30, color, 0.05)
        add_text(slide, body, x, 3.53, 3.40, 1.43, 15, INK)
        if i < 2:
            add_text(slide, "→", x + 3.51, 3.00, 0.55, 0.55, 27, LINE, True,
                     align=PP_ALIGN.CENTER)
    add_box(slide, 0.66, 5.56, 12.02, 0.79, LIGHT_BLUE, RGBColor(182, 204, 235))
    add_text(slide, "统一边界", 0.92, 5.82, 1.10, 0.31, 14, FUDAN, True)
    add_text(slide, "冻结或少量适配基础模型，在模型外构建可审计、可恢复、可部署的具身智能体运行系统。",
             2.08, 5.78, 10.15, 0.39, 16, INK, True)

    # 4. Prior work task definition.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agentic RAG-VLM 的研究问题与任务定义", 4, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_box(slide, 0.66, 1.42, 3.06, 4.94, PANEL, LINE)
    add_text(slide, "机器人抓取中的三类缺口", 0.93, 1.73, 2.55, 0.43, 17, FUDAN, True)
    add_bullets(slide, [
        "视觉相似不等于操作属性相容",
        "遮挡、支撑和易碎物形成场景约束",
        "开环规划无法诊断失败并持续改进",
    ], 0.93, 2.44, 2.54, 2.25, 14, FUDAN, 10)
    add_text(slide, "输入", 4.28, 1.58, 1.10, 0.34, 12, FUDAN, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 4.02, 2.06, 1.68, 1.68, LIGHT_BLUE, FUDAN)
    add_text(slide, "RGB-D\n+\n语言指令", 4.18, 2.42, 1.36, 0.92, 17, INK, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "→", 5.78, 2.60, 0.62, 0.50, 28, FUDAN, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "Agentic 推理", 6.42, 1.58, 2.10, 0.34, 12, TEAL, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 6.25, 2.06, 2.45, 1.68, LIGHT_TEAL, TEAL)
    add_text(slide, "经验检索 → 约束\n规划 → 评价\n反思 → 重试", 6.48, 2.37, 1.98, 1.07,
             15, INK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "→", 8.84, 2.60, 0.62, 0.50, 28, FUDAN, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "输出", 9.53, 1.58, 1.10, 0.34, 12, ORANGE, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 9.32, 2.06, 2.80, 1.68, LIGHT_ORANGE, ORANGE)
    add_text(slide, "g = (p, d, w, f, τ)", 9.48, 2.42, 2.46, 0.58, 19, INK, True,
             align=PP_ALIGN.CENTER, font="Cambria")
    add_text(slide, "位置 / 方向 / 宽度 / 力 / 抓取类型", 9.44, 3.12, 2.54, 0.30,
             11.5, MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 4.02, 4.52, 8.10, 1.19, WHITE, LINE)
    add_text(slide, "系统定位", 4.30, 4.83, 1.10, 0.32, 14, FUDAN, True)
    add_text(slide, "利用多模态大模型完成高层操作决策，再通过结构化参数连接机器人技能。",
             5.42, 4.76, 6.27, 0.52, 16, INK, True)

    # 5. Prior framework.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agentic RAG-VLM 框架总览：面向机器人抓取的结构化操作推理", 5, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_image_panel(slide, RAG_FIG / "fig.1.png", 4.20, 1.30, 8.44, 3.82, 0.06)
    explain = [
        ("任务输入", "RGB-D 图像、语言指令、目标物体；需要从感知结果转化为可执行抓取。", FUDAN),
        ("Agentic 外层", "检索操作经验、约束空间关系；失败后诊断、重试并写入记忆。", TEAL),
        ("结构化输出", "输出 g=(p,d,w,f,τ)，对应位置、方向、宽度、力和抓取类型。", ORANGE),
        ("核心定位", "不重新训练基座模型，而是在 VLM 外层组织推理、记忆和恢复流程。", BLUE),
    ]
    for i, (title, body, color) in enumerate(explain):
        y = 1.30 + i * 1.00
        add_box(slide, 0.64, y, 3.12, 0.78, WHITE, RGBColor(210, 218, 228), 0.8)
        add_rule(slide, 0.64, y, 3.12, color, 0.05)
        add_text(slide, title, 0.86, y + 0.18, 0.95, 0.24, 12.5, color, True)
        add_text(slide, body, 1.85, y + 0.13, 1.58, 0.39, 9.5, INK, True)
    # 6. HAA-RAG.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "HAA-RAG：按操作属性检索可复用经验", 6, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_image_panel(slide, RAG_FIG / "fig.2.png", 0.60, 1.40, 8.12, 4.98)
    add_box(slide, 9.02, 1.42, 3.25, 4.88, PANEL, LINE)
    add_text(slide, "全称", 9.32, 1.70, 1.70, 0.30, 15.5, FUDAN, True)
    add_text(slide, "Hierarchical Affordance-Aware\nRetrieval-Augmented Generation",
             9.32, 2.08, 2.50, 0.58, 10.2, INK, True)
    add_rule(slide, 9.32, 2.94, 2.36, FUDAN, 0.04)
    add_text(slide, "要解决的问题", 9.32, 3.17, 1.70, 0.30, 14.5, FUDAN, True)
    add_text(slide, "视觉相似的物体不一定能用相同抓取方式；需要检索“操作属性相容”的经验。",
             9.32, 3.55, 2.42, 0.62, 10.2, INK, True)
    add_rule(slide, 9.32, 4.30, 2.36, FUDAN, 0.04)
    add_text(slide, "检索流程", 9.32, 4.53, 1.30, 0.30, 14.5, FUDAN, True)
    add_bullets(slide, [
        "类别筛选：限定候选经验范围",
        "属性匹配：比较类型、材质、脆弱度和可抓取区域",
        "视觉重排：结合当前图像选择最相关经验",
    ], 9.30, 4.92, 2.55, 1.12, 9.5, FUDAN, 3)
    # 7. Scene graph.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Scene Graph：结构化场景状态与操作约束", 7, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_image_panel(slide, RAG_FIG / "fig.3.png", 0.60, 1.38, 8.54, 4.92)
    add_box(slide, 9.43, 1.42, 2.95, 4.88, PANEL, LINE)
    add_text(slide, "场景图表示什么", 9.70, 1.74, 1.88, 0.32, 16, TEAL, True)
    add_text(slide, "从 VLM 场景分析中构建对象节点与空间关系边，显式表示位置、类别、状态和脆弱度。",
             9.70, 2.24, 2.16, 0.78, 11.4, INK, True)
    add_rule(slide, 9.70, 3.20, 2.10, TEAL, 0.04)
    add_text(slide, "如何用于操作", 9.70, 3.45, 1.20, 0.30, 15, TEAL, True)
    add_bullets(slide, [
        "场景图构建：对象节点 + 空间关系边",
        "约束推断：状态保持 / 碰撞避免 / 支撑依赖 / 遮挡处理",
        "输出抓取参数调整：方向、夹持力、接近高度",
    ], 9.68, 3.88, 2.25, 1.30, 10.6, TEAL, 5)
    # 8. Reflection and memory.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Self-Reflection：失败诊断、分级恢复与情景记忆", 7, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_image_panel(slide, RAG_FIG / "fig.4.png", 0.58, 1.34, 8.33, 5.20)
    add_box(slide, 9.28, 1.48, 3.37, 4.93, PANEL, LINE)
    add_text(slide, "失败诊断与恢复", 9.60, 1.79, 2.02, 0.39, 18, FUDAN, True)
    add_text(slide, "14类失败分为4组，诊断后进入 L1-L3 分级恢复。",
             9.60, 2.22, 2.55, 0.42, 10.2, INK, True)
    retry = [
        ("L1", "调整参数", "夹持力、宽度、接近高度", BLUE),
        ("L2", "切换抓取类型", "Power / Pinch / Side", ORANGE),
        ("L3", "重新规划", "重置观察并生成新方案", RED),
    ]
    for i, (lv, title, desc, color) in enumerate(retry):
        y = 2.82 + i * 0.82
        add_text(slide, lv, 9.60, y, 0.51, 0.31, 14, color, True)
        add_text(slide, title, 10.20, y, 1.68, 0.31, 14, INK, True)
        add_text(slide, desc, 10.20, y + 0.35, 2.05, 0.29, 10.5, MUTED)
    add_rule(slide, 9.60, 5.35, 2.56, TEAL, 0.04)
    add_text(slide, "成功策略进入 Episodic Memory，\n后续同类对象可检索复用。",
             9.60, 5.60, 2.63, 0.57, 12, GREEN, True)

    # 8. Accepted paper and validation scope (detailed analytical rates stay off the main deck).
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agentic RAG-VLM 论文成果（IROS 2026 录用）", 8, "已有成果")
    add_footer(slide, "已有成果 · Agentic RAG-VLM")
    add_box(slide, 0.63, 1.29, 4.08, 5.57, PANEL, LINE)
    add_image_contain(slide, MIDTERM_ASSET / "agentic_rag_vlm_page1.png",
                      0.76, 1.42, 3.82, 5.29)
    add_label(slide, "IROS 2026 录用", 5.12, 1.41, 1.65, FUDAN)
    add_text(slide, "Agentic RAG-VLM", 5.12, 1.98, 5.92, 0.48, 22, FUDAN, True)
    add_text(slide, "Affordance-Aware Retrieval-Augmented Generation\nwith Self-Reflective Planning for Robotic Grasping",
             5.12, 2.48, 6.92, 0.73, 14, MUTED, True)
    add_text(slide, "主要技术贡献", 5.12, 3.49, 1.72, 0.36, 17, INK, True)
    add_bullets(slide, [
        "HAA-RAG：按操作属性检索抓取经验",
        "Scene Graph：将空间关系转化为参数约束",
        "Self-Reflection：结构化诊断、分级重试与记忆",
    ], 5.12, 4.02, 6.88, 1.27, 14, FUDAN, 8)
    add_box(slide, 5.12, 5.52, 7.20, 1.08, LIGHT_BLUE, RGBColor(182, 204, 235))
    add_text(slide, "验证范围", 5.40, 5.80, 1.06, 0.31, 14, FUDAN, True)
    add_text(slide, "论文采用 12 项任务的解析仿真进行组件分析，用于验证方法机制与消融趋势；\n物理执行证据由后续 Agentic VLA 的 MuJoCo 闭环实验补充。",
             6.49, 5.70, 5.50, 0.63, 13, INK, True)

    # 9. Research continuity.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "研究承接：Agentic RAG-VLM 与 Agentic VLA", 9, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA")
    add_box(slide, 0.64, 1.34, 5.35, 2.84, LIGHT_BLUE, RGBColor(194, 211, 238))
    add_label(slide, "已有成果", 0.92, 1.66, 1.20, FUDAN)
    add_text(slide, "Agentic RAG-VLM", 0.92, 2.18, 4.50, 0.36, 20, FUDAN, True)
    add_bullets(slide, ["用检索增强获取操作经验", "用场景图表示物体关系与约束", "用反思和记忆支持失败重试"],
                0.92, 2.76, 4.45, 1.18, 14, FUDAN, 6)

    add_text(slide, "→", 6.24, 2.36, 0.72, 0.70, 28, LINE, True,
             align=PP_ALIGN.CENTER)

    add_box(slide, 7.02, 1.34, 5.35, 2.84, LIGHT_TEAL, RGBColor(177, 220, 213))
    add_label(slide, "近期进展", 7.30, 1.66, 1.20, TEAL)
    add_text(slide, "Agentic VLA", 7.30, 2.18, 4.50, 0.36, 20, TEAL, True)
    add_bullets(slide, ["保留 Agentic 的监测、反思和恢复思想", "接入冻结 VLA 生成连续动作块", "加入 Runtime 管理推理时延和调用预算"],
                7.30, 2.76, 4.45, 1.18, 14, TEAL, 6)

    rows = [
        ("任务层面", "从桌面抓取规划，扩展到 VLA 控制下的连续操作流程"),
        ("机制层面", "从 RAG / Scene Graph / Reflection，延续到 Harness / Memory / Critic / Recovery"),
        ("模型接口", "从 VLM 输出高层操作参数，扩展到 VLA 输出可执行动作序列"),
        ("工程目标", "从提升决策可靠性，进一步关注闭环稳定性与高效推理"),
    ]
    y = 4.55
    for idx, (name, desc) in enumerate(rows):
        yy = y + idx * 0.52
        add_box(slide, 0.84, yy, 11.28, 0.43, WHITE, RGBColor(219, 225, 233), 0.6)
        add_text(slide, name, 1.07, yy + 0.08, 1.20, 0.24, 11.2,
                 FUDAN if idx < 2 else TEAL, True)
        add_text(slide, desc, 2.42, yy + 0.075, 9.28, 0.25, 11.2, INK)

    # 10. Agentic VLA overall framework.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agentic VLA：VLA 上的过程控制与闭环监督", 10, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA")
    add_image_panel(slide, AGENTIC_VLA_FIG / "fig1_framework.png",
                    4.02, 1.28, 8.64, 4.04, 0.06)
    explain = [
        ("Frozen VLA", "VLA 保持冻结，按需生成连续 action chunk。", BLUE),
        ("Process-Control Harness", "Risk Monitor 高频检测；VLM Critic 低频语义判断。", ORANGE),
        ("Optimize Runtime", "时延感知调度、动作复用和调用预算管理。", TEAL),
        ("研究定位", "系统层增强 VLA；语义判断低频接入。", FUDAN),
    ]
    for i, (title, body, color) in enumerate(explain):
        y = 1.30 + i * 1.00
        add_box(slide, 0.64, y, 3.02, 0.78, WHITE, RGBColor(210, 218, 228), 0.8)
        add_rule(slide, 0.64, y, 3.02, color, 0.05)
        add_text(slide, title, 0.84, y + 0.18, 1.38, 0.26, 11.2, color, True)
        add_text(slide, body, 2.26, y + 0.13, 1.08, 0.42, 9.2, INK, True)

    # 11. Execution supervision.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agentic Harness：执行监督与可重试控制循环", 11, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA")
    add_image_panel(slide, AGENTIC_VLA_FIG / "fig2_execution_supervision.png", 0.58, 1.44, 8.62, 4.98)
    add_box(slide, 9.48, 1.49, 3.15, 4.99, PANEL, LINE)
    add_text(slide, "一次控制周期", 9.78, 1.80, 2.30, 0.38, 17, FUDAN, True)
    add_bullets(slide, [
        "Observe：读取图像、状态与动作年龄",
        "Query：按需向冻结 VLA 请求动作块",
        "Execute：执行部分动作并监测进展",
        "Critic：识别停滞、陈旧或越界",
        "Recover：物理恢复、重观察与重规划",
        "Verify：验证结果，否则安全停止",
    ], 9.76, 2.32, 2.53, 3.38, 12.5, FUDAN, 6)
    # 12. Harness modules and recovery evidence.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Harness 内部机制：状态转换、情景记忆与有界恢复", 12, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA")
    add_image_panel(slide, AGENTIC_VLA_FIG / "fig3_modules.png", 0.58, 1.42, 7.24, 3.88)
    add_box(slide, 8.13, 1.47, 4.51, 3.70, PANEL, LINE)
    add_text(slide, "VLA Recovery Challenge", 8.44, 1.77, 3.70, 0.35, 16, FUDAN, True)
    headers = ["状态", "Continue", "Replan", "Retry", "Recovery"]
    rows = [
        ("T6 stall", ["PASS", "PASS", "PASS", "PASS"]),
        ("T9 stall", ["PASS", "PASS", "PASS", "PASS"]),
        ("T8 stale", ["FAIL", "FAIL", "FAIL", "STOP"]),
    ]
    xs = [8.43, 9.21, 10.03, 10.84, 11.63]
    ws = [0.72, 0.76, 0.76, 0.73, 0.72]
    for j, header in enumerate(headers):
        add_text(slide, header, xs[j], 2.34, ws[j], 0.27, 9, MUTED, True,
                 align=PP_ALIGN.CENTER)
    for i, (state, vals) in enumerate(rows):
        y = 2.82 + i * 0.66
        add_text(slide, state, xs[0], y + 0.10, ws[0], 0.27, 9.5, INK, True,
                 align=PP_ALIGN.CENTER)
        for j, val in enumerate(vals, 1):
            fill = GREEN if val == "PASS" else (ORANGE if val == "STOP" else RED)
            add_box(slide, xs[j], y, ws[j], 0.46, fill, fill, 0)
            add_text(slide, val, xs[j], y + 0.08, ws[j], 0.24, 8.5, WHITE, True,
                     align=PP_ALIGN.CENTER)
    add_text(slide, "精确状态恢复；不支持事件采用 fail-closed 安全停止。",
             8.44, 4.74, 3.67, 0.28, 9.5, MUTED, True)
    add_box(slide, 0.65, 5.48, 12.02, 0.89, PANEL, LINE)
    add_stat(slide, "2/2", "支持状态恢复并通过验证", 0.92, 5.58, 2.67, GREEN)
    add_stat(slide, "1", "不支持事件执行前安全停止", 3.93, 5.58, 2.67, ORANGE)
    add_stat(slide, "4.0–4.5×", "频繁重规划/重试的调用开销", 6.95, 5.58, 2.67, RED)
    add_stat(slide, "12", "每次物理恢复的有界动作数", 9.96, 5.58, 2.40, FUDAN)

    # 13. Optimize Runtime.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Optimize Runtime：面向实时约束的 VLA 调用管理", 13, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA")
    add_image_panel(slide, AGENTIC_VLA_FIG / "fig4_realtime_runtime.png", 0.56, 1.36, 8.86, 4.92)
    add_box(slide, 9.69, 1.36, 2.95, 4.92, PANEL, LINE)
    add_text(slide, "运行时机制", 9.98, 1.67, 2.36, 0.36, 16.5, TEAL, True)
    runtime_points = [
        ("1", "Low-step sampler", "少步采样，降低单次调用成本", FUDAN),
        ("2", "Call scheduler", "判断是否需要昂贵 full VLA call", TEAL),
        ("3", "Action reuse / cache", "未触发异常时复用动作块", ORANGE),
        ("4", "Runtime metrics", "记录调用次数、时延和实时超时", GREEN),
    ]
    for i, (idx, title, desc, color) in enumerate(runtime_points):
        y = 2.21 + i * 0.78
        add_box(slide, 9.97, y, 0.36, 0.36, color, color, 0)
        add_text(slide, idx, 10.02, y + 0.06, 0.25, 0.20, 9.5, WHITE, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, title, 10.45, y - 0.02, 1.66, 0.24, 11.4, color, True)
        add_text(slide, desc, 10.45, y + 0.29, 1.73, 0.29, 9.2, MUTED, True)
    add_rule(slide, 9.98, 5.55, 2.20, TEAL, 0.04)
    add_text(slide, "系统作用", 9.98, 5.72, 0.88, 0.24, 11.5, TEAL, True)
    add_text(slide, "Runtime 不替代 VLA，而是在执行循环中决定何时调用、复用、恢复和记录。",
             10.89, 5.65, 1.32, 0.42, 8.8, INK, True)

    # 14. Runtime and coupled results with editable charts.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "运行时优化证据：恢复结果与关键路径时延", 14, "近期进展")
    add_footer(slide, "近期进展 · Agentic VLA 联合证据")
    latency = ChartData()
    latency.categories = ["Eager BF16", "Compiled BF16", "Compiled + SMVE"]
    latency.add_series("P50", (155.88, 62.54, 52.54))
    latency.add_series("P95", (166.26, 65.75, 54.50))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.62), Inches(1.49),
        Inches(6.35), Inches(3.72), latency
    ).chart
    chart.has_legend = True
    chart.legend.font.size = Pt(9)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 180
    chart.value_axis.major_unit = 40
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(9)
    for idx, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = BLUE if idx == 0 else FUDAN
        series.format.line.fill.background()
    add_text(slide, "同状态恢复时延（ms）", 2.20, 1.25, 3.20, 0.29, 12, INK, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 7.29, 1.49, 5.34, 3.72, PANEL, LINE)
    add_text(slide, "80 ms policy-call deadline", 7.64, 1.82, 4.62, 0.34,
             16, FUDAN, True, align=PP_ALIGN.CENTER)
    add_stat(slide, "100%", "Eager BF16 miss rate", 7.62, 2.59, 1.48, RED)
    add_stat(slide, "0%", "Compiled BF16", 9.16, 2.59, 1.48, GREEN)
    add_stat(slide, "0%", "Compiled + SMVE", 10.70, 2.59, 1.57, GREEN)
    add_text(slide, "三种 profile 在两个精确恢复状态上均保持 2/2。",
             7.69, 4.36, 4.54, 0.39, 12.5, INK, True,
             align=PP_ALIGN.CENTER)
    add_box(slide, 0.64, 5.61, 12.02, 0.83, LIGHT_TEAL, RGBColor(170, 218, 207))
    add_text(slide, "结论", 0.93, 5.86, 0.72, 0.31, 14, TEAL, True)
    add_text(slide, "在相同恢复结果下，Compile / SMVE 相对 Eager 将 P95 降低 60.5% / 67.2%。",
             1.69, 5.81, 10.46, 0.42, 15, INK, True)

    # 15. Summary and plan.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "阶段总结与后续计划", 15, "总结")
    add_footer(slide)
    columns = [
        ("已有成果", FUDAN, [
            "Agentic RAG-VLM 获 IROS 2026 录用",
            "完成经验检索、场景约束、反思与记忆框架",
            "搭建 Agentic VLA Harness 与冻结 VLA 接入",
            "完成物理恢复、安全停止、运行轨迹与 Runtime 指标整理",
        ]),
        ("后续计划", TEAL, [
            "完善 VLM Critic / Planner 的事件触发协作",
            "深化 Runtime 推理优化：调用调度、动作复用与实时超时分析",
            "扩大物理仿真任务与长程闭环验证",
            "整理论文实验、系统图示与汇报材料",
        ]),
    ]
    for i, (title, color, bullets) in enumerate(columns):
        x = 0.63 + i * 6.28
        add_rule(slide, x, 1.50, 5.55, color, 0.075)
        add_text(slide, title, x, 1.83, 5.55, 0.47, 21, color, True)
        add_bullets(slide, bullets, x, 2.58, 5.45, 2.68, 13.2, color, 10)
    add_box(slide, 0.63, 5.47, 12.08, 0.92, LIGHT_BLUE, RGBColor(182, 204, 235))
    add_text(slide, "研究主线", 0.93, 5.78, 1.10, 0.31, 14, FUDAN, True)
    add_text(slide, "让多模态模型不仅能理解任务，还能在执行过程中持续监测、恢复并满足实时部署约束。",
             2.08, 5.73, 10.10, 0.43, 16, INK, True)
    # 16. Thanks.
    slide = prs.slides.add_slide(blank)
    add_box(slide, 0, 0, 13.333, 7.5, FUDAN, FUDAN, 0)
    add_text(slide, "谢谢各位老师和同学", 1.10, 2.55, 11.10, 0.70, 34, WHITE, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "敬请批评指正", 1.10, 3.42, 11.10, 0.44, 20, RGBColor(226, 232, 240), True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "陈涛 · 24210860031", 1.10, 4.55, 11.10, 0.32, 14, RGBColor(226, 232, 240), False,
             align=PP_ALIGN.CENTER)

    # 17. Backup: Agentic RAG-VLM validation overview.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "备份：Agentic RAG-VLM 实验验证概览", 16, "备份")
    add_footer(slide, "备份 · Agentic RAG-VLM")
    add_text(slide, "Component Analysis on 12-task Benchmark", 0.68, 1.34, 7.25, 0.34, 15, FUDAN, True)
    headers = ["Config.", "SG", "IT", "LH", "Overall", "Delta"]
    rows = [
        ["Ours", "91.7", "64.2", "66.7", "78.3", "--"],
        ["w/o HAA-RAG", "50.0", "0.0", "0.0", "25.0", "-53.3"],
        ["w/o Scene Graph", "91.7", "25.0", "66.7", "65.3", "-13.0"],
        ["w/o Recovery", "63.3", "33.3", "23.3", "46.7", "-31.6"],
        ["VLM-Only", "50.0", "0.0", "0.0", "25.0", "-53.3"],
    ]
    col_x = [0.70, 3.08, 3.95, 4.82, 5.69, 6.70]
    col_w = [2.28, 0.76, 0.76, 0.76, 0.92, 0.82]
    y0, row_h = 1.83, 0.49
    for j, (h, x, w) in enumerate(zip(headers, col_x, col_w)):
        add_box(slide, x, y0, w, row_h, LIGHT_BLUE, RGBColor(182, 204, 235), 0.5)
        add_text(slide, h, x + 0.03, y0 + 0.13, w - 0.06, 0.20, 9.6, FUDAN, True,
                 align=PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT)
    for i, row in enumerate(rows):
        y = y0 + (i + 1) * row_h
        fill = WHITE if i % 2 == 0 else PANEL
        for j, (cell, x, w) in enumerate(zip(row, col_x, col_w)):
            add_box(slide, x, y, w, row_h, fill, RGBColor(218, 225, 233), 0.45)
            color = FUDAN if i == 0 else INK
            add_text(slide, cell, x + 0.03, y + 0.13, w - 0.06, 0.20, 9.5, color, i == 0,
                     align=PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT)
    add_text(slide, "SG/IT/LH: single-grasp / interactive / long-horizon.\n12 tasks x 30 trials; values are success rates (%); Delta: drop vs. full system.",
             0.72, 5.01, 7.60, 0.45, 8.8, MUTED)

    add_rule(slide, 8.50, 1.50, 3.55, TEAL, 0.075)
    add_text(slide, "机制结论", 8.50, 1.83, 3.55, 0.47, 21, TEAL, True)
    add_bullets(slide, [
        "HAA-RAG 是经验检索基础；移除后整体成功率降至 25.0%",
        "Scene Graph 主要影响交互约束任务",
        "Self-Reflection / Recovery 显著影响失败恢复和长程稳定性",
    ], 8.50, 2.58, 3.55, 2.44, 12.5, TEAL, 10)
    add_box(slide, 0.63, 5.47, 12.08, 0.92, LIGHT_TEAL, RGBColor(170, 218, 207))
    add_text(slide, "使用方式", 0.93, 5.78, 1.10, 0.31, 14, TEAL, True)
    add_text(slide, "主线不主动展开具体数值；如被问到实验验证，用本页说明消融表和机制结论。",
             2.08, 5.73, 10.10, 0.43, 15.2, INK, True)

    # 17. Backup: execution supervision.
    slide = prs.slides.add_slide(blank)
    add_header(slide, "备份：Agentic Harness 执行监督与可重试控制循环", 17, "备份")
    add_footer(slide, "备份 · Agentic VLA")
    add_image_panel(slide, AGENTIC_VLA_FIG / "fig2_execution_supervision.png", 0.58, 1.44, 8.62, 4.98)
    add_box(slide, 9.48, 1.49, 3.15, 4.99, PANEL, LINE)
    add_text(slide, "一次控制周期", 9.78, 1.80, 2.30, 0.38, 17, FUDAN, True)
    add_bullets(slide, [
        "Observe：读取图像、状态与动作年龄",
        "Query：按需向冻结 VLA 请求动作块",
        "Execute：执行部分动作并监测进展",
        "Critic：识别停滞、陈旧或越界",
        "Recover：物理恢复、重观察与重规划",
        "Verify：验证结果，否则安全停止",
    ], 9.76, 2.32, 2.53, 3.38, 12.5, FUDAN, 6)

    # Midterm detailed version: merge the old task-definition page into the
    # Agentic RAG-VLM overview. Keep selected mechanism pages as backup, but
    # hide result-heavy or redundant summary pages from the main projected deck.
    for slide_index in sorted([2, 3, 8, 11, 12, 14], reverse=True):
        remove_slide(prs, slide_index)
    renumber_slides(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    print(f"slides={len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
